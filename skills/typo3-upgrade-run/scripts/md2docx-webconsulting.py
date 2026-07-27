#!/usr/bin/env python3
"""
Markdown -> DOCX in the webconsulting corporate design.

Design tokens come from skills/webconsulting-branding:
  primary #1b7a95 · primary-dark #155d73 · primary-strong #0f4555
  accent  #66c4e1 · pale #e8f4f8 · ink #171a1d · muted #5e6870
  Typography: Hanken Grotesk. Surfaces: borderless and square.

The markdown file stays the single source of truth; this only re-renders it.
"""
import re
import sys
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

PRIMARY        = RGBColor(0x1b, 0x7a, 0x95)
PRIMARY_DARK   = RGBColor(0x15, 0x5d, 0x73)
PRIMARY_STRONG = RGBColor(0x0f, 0x45, 0x55)
ACCENT_HEX     = "66c4e1"
PALE_HEX       = "e8f4f8"
INK            = RGBColor(0x17, 0x1a, 0x1d)
MUTED          = RGBColor(0x5e, 0x68, 0x70)
FONT           = "Hanken Grotesk"
FONT_FALLBACK  = "Segoe UI"


def shade(cell, hex_fill):
    """Flat fill, no border — the CI uses borderless square surfaces."""
    tcPr = cell._tc.get_or_add_tcPr()
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:color"), "auto")
    el.set(qn("w:fill"), hex_fill)
    tcPr.append(el)


def clear_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "nil")
        borders.append(e)
    tblPr.append(borders)


def rule(paragraph, hex_color, size=6):
    """A square accent rule under a heading."""
    pPr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), hex_color)
    bdr.append(bottom)
    pPr.append(bdr)


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def add_runs(par, text, base_size=10, color=INK, bold=False):
    """Render **bold**, `code` and [link](target) as runs."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            # Inline spans nest: **`code`** is common. Render the inner code span as
            # a bold code run instead of leaking its backticks into the text.
            inner = part[2:-2]
            for sub in re.split(r"(`[^`]+`)", inner):
                if not sub:
                    continue
                if sub.startswith("`") and sub.endswith("`"):
                    r = par.add_run(sub[1:-1])
                    r.font.name = "Consolas"
                    r.font.color.rgb = PRIMARY_DARK
                    r.font.size = Pt(base_size - 0.5)
                else:
                    r = par.add_run(sub)
                    r.font.name = FONT
                    r.font.size = Pt(base_size)
                    r.font.color.rgb = color
                r.bold = True
            continue
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.color.rgb = PRIMARY_DARK
            r.font.size = Pt(base_size - 0.5)
            continue
        elif part.startswith("[") and "](" in part:
            # Link labels frequently wrap a code span: [`file.md`](path). Word has no
            # nested-run concept here, so render the label as plain link text.
            label = part[1 : part.index("]")].replace("`", "").replace("**", "")
            r = par.add_run(label)
            r.font.color.rgb = PRIMARY
            r.underline = True
        else:
            r = par.add_run(re.sub(r"\\([_*])", r"\1", part))
            r.bold = bold
        r.font.name = FONT
        r.font.size = Pt(base_size)
        if r.font.color.rgb is None:
            r.font.color.rgb = color


def build(md_path, out_path, subtitle=None):
    lines = open(md_path, encoding="utf-8").read().split("\n")
    doc = Document()

    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_FALLBACK)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.15

    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.2)
        s.top_margin = Cm(2.0)
        s.bottom_margin = Cm(1.8)
        # header-only brand notice, per the CI rules
        hp = s.header.paragraphs[0]
        hr = hp.add_run("web:consulting")
        hr.font.name = FONT
        hr.font.size = Pt(9)
        hr.font.bold = True
        hr.font.color.rgb = PRIMARY
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        # footer stays quiet: no border, no repeated notice
        fp = s.footer.paragraphs[0]
        fr = fp.add_run("Interner Bericht · nur lokale DDEV-Messung")
        fr.font.name = FONT
        fr.font.size = Pt(8)
        fr.font.color.rgb = MUTED
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT

    i = 0
    in_table = False
    table_rows = []

    def flush_table():
        nonlocal table_rows
        if not table_rows:
            return
        cols = len(table_rows[0])
        t = doc.add_table(rows=0, cols=cols)
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        clear_table_borders(t)
        for ri, row in enumerate(table_rows):
            cells = t.add_row().cells
            for ci in range(cols):
                cell = cells[ci]
                cell.text = ""
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                val = row[ci] if ci < len(row) else ""
                add_runs(p, val, base_size=9,
                         color=RGBColor(0xff, 0xff, 0xff) if ri == 0 else INK,
                         bold=(ri == 0))
                if ri == 0:
                    shade(cell, "1b7a95")
                elif ri % 2 == 0:
                    shade(cell, PALE_HEX)
        doc.add_paragraph()
        table_rows = []

    while i < len(lines):
        ln = lines[i]

        if ln.strip().startswith("```"):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i]); i += 1
            i += 1
            for b in buf:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.left_indent = Cm(0.5)
                r = p.add_run(b)
                r.font.name = "Consolas"
                r.font.size = Pt(8.5)
                r.font.color.rgb = PRIMARY_STRONG
            doc.add_paragraph()
            continue

        if ln.strip().startswith("|"):
            cells = [c.strip() for c in ln.strip().strip("|").split("|")]
            if all(set(c) <= set("-: ") for c in cells if c):
                i += 1; continue
            table_rows.append(cells)
            in_table = True
            i += 1
            continue
        elif in_table:
            flush_table(); in_table = False

        if ln.startswith("# "):
            p = doc.add_paragraph()
            add_runs(p, ln[2:], base_size=22, color=PRIMARY_STRONG, bold=True)
            p.paragraph_format.space_after = Pt(2)
            rule(p, ACCENT_HEX, size=12)
            if subtitle:
                sp = doc.add_paragraph()
                add_runs(sp, subtitle, base_size=10, color=MUTED)
                subtitle = None
        elif ln.startswith("## "):
            doc.add_paragraph()
            p = doc.add_paragraph()
            add_runs(p, ln[3:], base_size=14, color=PRIMARY, bold=True)
            p.paragraph_format.space_after = Pt(3)
            rule(p, PALE_HEX, size=6)
        elif ln.startswith("### "):
            p = doc.add_paragraph()
            add_runs(p, ln[4:], base_size=11, color=PRIMARY_DARK, bold=True)
        elif ln.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.4)
            add_runs(p, ln[2:], base_size=9.5, color=MUTED)
        elif re.match(r"^[-*] |^\d+\. ", ln):
            ordered = bool(re.match(r"^\d+\. ", ln))
            body = re.sub(r"^([-*] |\d+\.\s*)", "", ln)
            # A list item continues onto indented following lines.
            while i + 1 < len(lines):
                nxt = lines[i + 1]
                if (not nxt.strip() or not nxt.startswith(("  ", "\t"))
                        or re.match(r"^\s*([-*] |\d+\. )", nxt)):
                    break
                body += " " + nxt.strip()
                i += 1
            p = doc.add_paragraph(style="List Number" if ordered else "List Bullet")
            add_runs(p, body)
        elif ln.strip() == "---":
            pass
        elif ln.strip() == "":
            pass
        else:
            # Markdown joins soft-wrapped lines into one paragraph. Rendering each
            # source line separately splits inline spans -- a **bold** phrase broken
            # across two lines would emit literal asterisks.
            buf = [ln]
            while i + 1 < len(lines):
                nxt = lines[i + 1]
                if (not nxt.strip() or nxt.startswith(("#", ">", "|", "```"))
                        or re.match(r"^([-*] |\d+\. )", nxt) or nxt.strip() == "---"):
                    break
                buf.append(nxt.strip())
                i += 1
            p = doc.add_paragraph()
            add_runs(p, " ".join(x.strip() for x in buf))
        i += 1

    flush_table()
    doc.save(out_path)
    print(f"written: {out_path}")


if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2],
          subtitle=sys.argv[3] if len(sys.argv) > 3 else None)
